"""Decks Compii writes: real PowerPoint files that look made, not generated.

The slides are drawn edge to edge rather than poured into placeholders: a 16:9 canvas,
one type scale, one accent, generous margins, and a page number on every slide after
the cover. Three themes keep a pitch, a lesson and a report from looking the same.

Slides come in as plain data::

    {"kind": "content", "title": "What changed", "bullets": ["...", "..."], "notes": "..."}
    {"kind": "section", "title": "Where we are", "subtitle": "Three months in"}
    {"kind": "quote", "text": "Sell the problem, not the product.", "by": "Someone"}
    {"kind": "stats", "title": "By the numbers", "stats": [{"value": "42%", "label": "faster"}]}
    {"kind": "closing", "title": "Thank you", "subtitle": "Questions?"}
"""

from __future__ import annotations

import math
from datetime import datetime
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = 13.333, 7.5
SLIDE_W, SLIDE_H = 12192000, 6858000  # exactly 16:9, where Inches(13.333) rounds off by 305 EMU
MARGIN = 0.9
FONT = "Segoe UI"

THEMES: dict[str, dict[str, str]] = {
    "ink": {"bg": "0E1526", "panel": "172138", "text": "F4F1EA", "muted": "A6AEC4",
            "accent": "7C8CFF", "line": "2A3550"},
    "paper": {"bg": "FBF9F5", "panel": "F1EDE5", "text": "17181C", "muted": "5B6070",
              "accent": "4F46E5", "line": "E2DCD1"},
    "sea": {"bg": "0B1D24", "panel": "12303B", "text": "F1F7F5", "muted": "9DB9BB",
            "accent": "5FD6C2", "line": "1E434E"},
}

MAX_SLIDES = 60


class DeckError(ValueError):
    pass


def deck_bytes(title: str, slides: list[dict[str, Any]], subtitle: str = "", author: str = "",
               theme: str = "ink", date: str | None = None) -> bytes:
    """Render a deck and return its .pptx bytes. Raises DeckError on empty input."""
    title = str(title or "").strip()[:140]
    subtitle = str(subtitle or "").strip()[:240]
    author = str(author or "").strip()[:80]
    if not title:
        raise DeckError("A deck needs a title.")
    clean = [s for s in slides if isinstance(s, dict)]
    if not clean:
        raise DeckError("A deck needs at least one slide besides the cover.")
    if len(clean) > MAX_SLIDES:
        raise DeckError(f"A deck fits at most {MAX_SLIDES} slides.")
    colors = THEMES.get(str(theme or "ink").lower(), THEMES["ink"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    when = date or datetime.now().strftime("%d %B %Y")
    blank = prs.slide_layouts[6]

    _cover(prs, blank, colors, title, subtitle, author, when)
    total = len(clean) + 1
    part = 0
    for n, spec in enumerate(clean, start=2):
        kind = str(spec.get("kind") or "content").lower()
        if kind == "section":
            part += 1
        slide = prs.slides.add_slide(blank)
        _paint(slide, colors, panel=kind in ("section", "closing"))
        if kind == "section":
            _section(slide, colors, spec, part)
        elif kind == "quote":
            _quote(slide, colors, spec)
        elif kind == "stats":
            _stats(slide, colors, spec)
        elif kind == "closing":
            _closing(slide, colors, spec)
        else:
            _content(slide, colors, spec)
        _footer(slide, colors, title, n, total)
        notes = str(spec.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes[:4000]

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- slide kinds


def _cover(prs, layout, c, title: str, subtitle: str, author: str, when: str) -> None:
    slide = prs.slides.add_slide(layout)
    _paint(slide, c)
    # A quiet accent block in the far corner keeps the cover from feeling empty.
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(W - 2.65), Inches(0), Inches(2.65), Inches(H))
    _fill(block, c["accent"], alpha_soft=True)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(W - 0.22), Inches(0), Inches(0.22), Inches(H))
    _fill(stripe, c["accent"])
    _bar(slide, c, MARGIN, 6.2, 1.6, 0.07)
    _text(slide, MARGIN, 6.34, 7.0, 0.4, when + (f"  ·  {author}" if author else ""), 12, c["muted"])
    _text(slide, MARGIN, 1.95, 9.9, 2.2, title, _cover_size(title), c["text"], bold=True, spacing=0.95)
    if subtitle:
        _text(slide, MARGIN, 4.15, 9.6, 1.3, subtitle, 21, c["muted"], spacing=1.2)


def _content(slide, c, spec: dict[str, Any]) -> None:
    title = _s(spec.get("title"), 140)
    bullets = _bullets(spec)
    _heading(slide, c, spec, title)
    if not bullets:
        body = _s(spec.get("text") or spec.get("body"), 900)
        if body:
            _text(slide, MARGIN, 2.15, 11.4, 4.2, body, 19, c["text"], spacing=1.3)
        return
    two_col = len(bullets) > 6 and len(bullets) <= 12
    size = _bullet_size(bullets, two_col)
    if two_col:
        half = math.ceil(len(bullets) / 2)
        _bullet_column(slide, c, bullets[:half], MARGIN, 2.05, 5.45, size)
        _bullet_column(slide, c, bullets[half:], MARGIN + 5.95, 2.05, 5.45, size)
    else:
        _bullet_column(slide, c, bullets[:10], MARGIN, 2.05, 11.4, size)


def _bullet_column(slide, c, bullets: list[str], x: float, y: float, width: float, size: int) -> None:
    top = y
    chars = max(24, int(width * 72 / (size * 0.52)))
    for text in bullets:
        lines = max(1, math.ceil(len(text) / chars))
        height = lines * (size * 1.28 / 72) + 0.12
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(top + size / 118), Inches(0.09), Inches(0.09))
        _fill(dot, c["accent"])
        _text(slide, x + 0.28, top, width - 0.28, height, text, size, c["text"], spacing=1.22)
        top += height + 0.16


def _section(slide, c, spec: dict[str, Any], part: int) -> None:
    _text(slide, MARGIN, 2.1, 4.0, 0.5, f"PART {part:02d}", 13, c["accent"], bold=True)
    _text(slide, MARGIN, 2.55, 11.2, 1.5, _s(spec.get("title"), 140) or "Next", 44, c["text"], bold=True, spacing=1.0)
    sub = _s(spec.get("subtitle") or spec.get("text"), 300)
    if sub:
        _text(slide, MARGIN, 4.15, 9.8, 1.2, sub, 20, c["muted"], spacing=1.25)
    _bar(slide, c, MARGIN, 3.95, 1.3, 0.06)


def _quote(slide, c, spec: dict[str, Any]) -> None:
    _text(slide, MARGIN - 0.06, 1.15, 2.0, 1.4, "“", 110, c["accent"], bold=True)
    quote = _s(spec.get("text") or spec.get("title"), 600)
    _text(slide, MARGIN, 2.45, 10.9, 2.5, quote, _quote_size(quote), c["text"], italic=True, spacing=1.18)
    by = _s(spec.get("by"), 120)
    if by:
        _bar(slide, c, MARGIN, 5.35, 0.8, 0.05)
        _text(slide, MARGIN, 5.55, 8.0, 0.5, by, 15, c["muted"], bold=True)


def _stats(slide, c, spec: dict[str, Any]) -> None:
    _heading(slide, c, spec, _s(spec.get("title"), 140) or "By the numbers")
    stats = [s for s in spec.get("stats") or [] if isinstance(s, dict)][:4]
    if not stats:
        return
    gap = 0.35
    width = (W - 2 * MARGIN - gap * (len(stats) - 1)) / len(stats)
    for i, stat in enumerate(stats):
        x = MARGIN + i * (width + gap)
        if i:
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - gap / 2), Inches(2.05),
                                          Inches(0.012), Inches(3.3))
            _fill(rule, c["line"])
        value = _s(stat.get("value"), 24) or "—"
        label = _s(stat.get("label"), 90)
        _text(slide, x, 2.1, width, 1.35, value, _stat_size(value), c["accent"], bold=True, spacing=0.95)
        _text(slide, x, 3.5, width, 1.6, label, 14, c["muted"], spacing=1.3)
        note = _s(stat.get("note"), 140)
        if note:
            _text(slide, x, 4.55, width, 0.9, note, 12, c["muted"], spacing=1.25)


def _closing(slide, c, spec: dict[str, Any]) -> None:
    _text(slide, MARGIN, 2.45, 11.2, 1.3, _s(spec.get("title"), 140) or "Thank you", 48, c["text"],
          bold=True, spacing=1.0)
    sub = _s(spec.get("subtitle") or spec.get("text"), 300)
    if sub:
        _text(slide, MARGIN, 3.85, 9.8, 1.2, sub, 20, c["muted"], spacing=1.3)
    _bar(slide, c, MARGIN, 3.62, 1.3, 0.06)


# ---------------------------------------------------------------- helpers


def _heading(slide, c, spec: dict[str, Any], title: str) -> None:
    eyebrow = _s(spec.get("eyebrow") or spec.get("label"), 60)
    top = 0.72
    if eyebrow:
        _text(slide, MARGIN, top, 8.0, 0.32, eyebrow.upper(), 11, c["accent"], bold=True)
        top += 0.38
    _text(slide, MARGIN, top, 11.5, 0.95, title or "…", _title_size(title), c["text"], bold=True, spacing=1.05)
    _bar(slide, c, MARGIN, top + 0.92, 1.15, 0.055)


def _footer(slide, c, deck_title: str, n: int, total: int) -> None:
    line_color = c["line"]
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(6.92),
                                  Inches(W - 2 * MARGIN), Inches(0.01))
    _fill(rule, line_color)
    _text(slide, MARGIN, 7.02, 8.2, 0.3, deck_title, 10, c["muted"])
    _text(slide, W - MARGIN - 1.2, 7.02, 1.2, 0.3, f"{n:02d} / {total:02d}", 10, c["muted"], align=PP_ALIGN.RIGHT)


def _paint(slide, c, *, panel: bool = False) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(c["panel"] if panel else c["bg"])


def _fill(shape, color: str, *, alpha_soft: bool = False) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if alpha_soft:
        # A slightly lifted panel rather than a flat block of accent.
        from pptx.oxml.ns import qn

        solid = shape.fill._xPr.find(qn("a:solidFill"))
        if solid is not None:
            srgb = solid.find(qn("a:srgbClr"))
            if srgb is not None:
                alpha = srgb.makeelement(qn("a:alpha"), {"val": "16000"})
                srgb.append(alpha)


def _bar(slide, c, x: float, y: float, w: float, h: float) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(shape, c["accent"])


def _text(slide, x: float, y: float, w: float, h: float, text: str, size: int, color: str, *,
          bold: bool = False, italic: bool = False, align=PP_ALIGN.LEFT, spacing: float = 1.15) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    run = paragraph.add_run()
    run.text = text
    font = run.font
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = _rgb(color)


def _s(value: Any, limit: int) -> str:
    return str(value or "").strip()[:limit]


def _bullets(spec: dict[str, Any]) -> list[str]:
    raw = spec.get("bullets")
    if raw is None:
        body = spec.get("body")
        raw = body.splitlines() if isinstance(body, str) else []
    return [_s(b, 400) for b in (raw or []) if _s(b, 400)][:12]


def _rgb(value: str) -> RGBColor:
    value = str(value).strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _cover_size(title: str) -> int:
    if len(title) > 58:
        return 40
    if len(title) > 34:
        return 46
    return 54


def _title_size(title: str) -> int:
    return 30 if len(title) <= 52 else 26


def _quote_size(quote: str) -> int:
    if len(quote) > 220:
        return 22
    if len(quote) > 120:
        return 26
    return 30


def _stat_size(value: str) -> int:
    return 44 if len(value) <= 7 else 34


def _bullet_size(bullets: list[str], two_col: bool) -> int:
    total = sum(len(b) for b in bullets)
    longest = max(len(b) for b in bullets)
    if two_col:
        return 18 if total <= 520 and longest <= 110 else 16
    if total <= 240 and longest <= 85:
        return 21
    if total <= 420 and longest <= 120:
        return 19
    return 17
