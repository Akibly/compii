"""Leads and marketing: the small business work Compii gets asked to do on a PC.

Only public, lawful information belongs here: details from a page the user is looking at,
a trade-show list, an export, or something they dictate. Compii saves them as a real file
the user can open, never sends anything by itself, and keeps the list free of duplicates.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any, Callable

from .policy import ActionKind
from .tools import Tool, ToolError, ToolResult

COLUMNS = ["name", "company", "role", "email", "phone", "website", "location", "source", "interest", "next_step",
           "status"]
MAX_LEADS = 400


def _key(lead: dict[str, Any]) -> str:
    """What makes a lead this lead: their email, else their phone, else company + name."""
    email = str(lead.get("email", "")).strip().lower()
    if email:
        return f"email:{email}"
    phone = re.sub(r"\D", "", str(lead.get("phone", "")))
    if len(phone) >= 7:
        return f"phone:{phone[-9:]}"
    return "name:" + (str(lead.get("company", "")).strip().lower() + "|" + str(lead.get("name", "")).strip().lower())


def _read(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            return [dict(r) for r in csv.DictReader(f)]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = [[("" if c.value is None else str(c.value)) for c in row] for row in ws.iter_rows(max_row=MAX_LEADS + 1)]
    wb.close()
    if not rows:
        return []
    head = [h.strip().lower().replace(" ", "_") for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(x.strip() for x in r)]


def _write(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        return
    from openpyxl import Workbook

    from .documents import style_worksheet

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([str(r.get(c, ""))[:200] for c in COLUMNS])
    style_worksheet(ws)
    wb.save(path)


def lead_tools(workspace: Path, resolve: Callable[[str], Path]) -> list[Tool]:
    s = {"type": "string"}

    def save_leads(a: dict[str, Any]) -> ToolResult:
        raw = [l for l in a["leads"] if isinstance(l, dict) and any(str(v).strip() for v in l.values())]
        if not raw:
            raise ToolError("No leads to save.")
        path = resolve(a.get("path") or f"leads/{a['list']}.csv")
        if path.suffix.lower() not in (".csv", ".xlsx"):
            path = path.with_suffix(".csv")
        rows = _read(path)
        by_key = {_key(r): i for i, r in enumerate(rows)}
        added, merged = 0, 0
        for lead in raw[:MAX_LEADS]:
            row = {c: str(lead.get(c, "")).strip()[:300] for c in COLUMNS}
            if not row["source"]:
                row["source"] = "from the user"
            k = _key(row)
            if k in by_key:
                old = rows[by_key[k]]
                for c in COLUMNS:
                    if row[c] and not old.get(c):
                        old[c] = row[c]
                        merged += 1
                continue
            by_key[k] = len(rows)
            rows.append(row)
            added += 1
        _write(path, rows)
        return ToolResult(
            f"Saved {added} new lead(s)" + (f", filled in {merged} detail(s) on ones you already had" if merged else "")
            + f" - the list now has {len(rows)} in {path}. Open it or ask me to work through it.",
            {"op": "created" if not path.exists() else "changed", "path": str(path)}, files=[(str(path), "created")])

    def list_leads(a: dict[str, Any]) -> ToolResult:
        folder = resolve(a.get("folder") or "leads")
        files = sorted(folder.glob("*.csv")) + sorted(folder.glob("*.xlsx")) if folder.is_dir() else []
        if not files:
            return ToolResult("No lead lists yet. Give me details from a page, a file or your own notes and I'll "
                              "start one.")
        out = []
        for f in files:
            rows = _read(f)
            open_ones = [r for r in rows if (r.get("status") or "").lower() not in ("won", "lost", "done", "closed")]
            due = [r for r in open_ones if (r.get("next_step") or "").strip()]
            out.append(f"- {f.name}: {len(rows)} leads, {len(open_ones)} still open, {len(due)} with a next step "
                       f"written down")
        return ToolResult("\n".join(out))

    def plan_marketing(a: dict[str, Any]) -> ToolResult:
        """Write the plan and calendar the model composed into real files, in one step."""
        from .documents import docx_bytes, tidy_markdown

        base = resolve(a["file"]).with_suffix("")
        md = base.with_suffix(".md")
        written = []
        if a.get("plan"):
            plan = tidy_markdown(str(a["plan"]))
            md.parent.mkdir(parents=True, exist_ok=True)
            md.write_text(plan, encoding="utf-8")
            written.append(str(md))
            # A plan gets shared with people, so it comes out as a Word file too.
            docx = base.with_suffix(".docx")
            docx.write_bytes(docx_bytes(plan, title=base.name.replace("-", " ").replace("_", " ")))
            written.append(str(docx))
        if a.get("calendar"):
            path = base.parent / f"{base.name}-calendar.xlsx"
            rows = [{"name": str(c.get("when", "")), "company": str(c.get("channel", "")),
                     "interest": str(c.get("idea", "")), "next_step": str(c.get("call_to_action", "")),
                     "status": "planned"} for c in a["calendar"] if isinstance(c, dict)]
            _write(path, rows)
            written.append(str(path))
        if not written:
            raise ToolError("Give me the plan text, or a calendar of posts, or both.")
        return ToolResult("Written: " + ", ".join(written) + ". Tell the user what's in them and where they are.",
                          files=[(p, "created") for p in written])

    lead_props = {"list": {**s, "description": "Which list, e.g. 'Bakery customers'."},
                  "leads": {"type": "array", "items": {"type": "object", "properties": {
                      c: s for c in COLUMNS}, "additionalProperties": False},
                      "description": "One entry per person or company."},
                  "path": {**s, "description": "Optional file; defaults to leads/<list>.csv"}}
    return [
        Tool("save_leads", "Add leads or customers to a list file (CSV or Excel), without duplicates: an existing "
             "email or phone is filled in rather than added again. Only store details the user gave you or read from "
             "a page they are looking at.",
             lead_props, ["list", "leads"], ActionKind.CREATE, save_leads,
             lambda a: f"Save {len(a['leads'])} lead(s) to {a['list']}"),
        Tool("list_leads", "The lead lists you have and how many are still open.",
             {"folder": {**s, "description": "Folder of lists; default is the leads folder in Compii's workspace."}},
             [], ActionKind.READ, list_leads, lambda a: "Look at the lead lists"),
        Tool("plan_marketing", "Write a marketing plan (Markdown and a shareable Word copy) and an optional content "
             "calendar (Excel) you have composed, in one step. It only saves the files; nothing is posted or sent.",
             {"file": {**s, "description": "Where to save, e.g. Documents/Marketing/bakery-launch."},
              "plan": {**s, "description": "The full plan in Markdown: audience, message, channels, offers, budget, "
                                          "how to measure it."},
              "calendar": {"type": "array", "items": {"type": "object", "properties": {
                  "when": s, "channel": s, "idea": s, "call_to_action": s}, "additionalProperties": False},
                  "description": "Dated posts or mailings: when, where, the idea, what you want the reader to do."}},
             ["file"], ActionKind.CREATE, plan_marketing,
             lambda a: f"Write the marketing plan {Path(a['file']).name}"),
    ]
