"""Follow-ups keep their context with Claude's real reply objects, and the honesty checks stop
turning ordinary answers into failures or wiping facts about work done earlier in the thread."""

import asyncio

from anthropic.types.beta import BetaServerToolUseBlock, BetaTextBlock, BetaToolUseBlock

from compii_core.agent import BrainReply, compact_history, unbacked_claims
from compii_core.keys import MemoryKeyStore
from compii_core.server import build_app

from test_core import home, quiet_voice  # noqa: F401 - fixture
from test_honest_and_coach import Scripted, _final


def _app(home, brain):
    return build_app(brain=brain, workspace=home["workspace"], roots=[home["downloads"]], token="t",
                     keys=MemoryKeyStore(), voice=quiet_voice(home["workspace"]))


def test_compact_history_reads_claude_sdk_blocks():
    msgs = [
        {"role": "user", "content": "make notes.txt"},
        {"role": "assistant", "content": [BetaToolUseBlock(id="a", name="write_text_file", input={}, type="tool_use")]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "a", "content": "ok"}]},
        {"role": "assistant", "content": [BetaTextBlock(text="Done - notes.txt is in Documents.", type="text")]},
    ]
    assert compact_history(msgs) == [{"role": "user", "content": "make notes.txt"},
                                     {"role": "assistant", "content": "(Compii ran: write_text_file)\nDone - notes.txt is in Documents."}]


class ClaudeShaped(Scripted):
    """Like Scripted, but replies with the SDK objects the real Claude brain returns."""

    async def step(self, system, messages, tools):
        reply = await super().step(system, messages, tools)
        blocks = [BetaTextBlock(text=b["text"], type="text") if b["type"] == "text" else
                  BetaToolUseBlock(id=b["id"], name=b["name"], input=b["input"], type="tool_use")
                  for b in reply.content]
        return BrainReply(blocks, reply.text, reply.tool_calls, reply.stop_reason)


def test_a_follow_up_with_claude_still_sees_the_earlier_answer(home):
    brain = ClaudeShaped(["Ghana uses GMT.", "Yes, the same as London in winter."])
    app = _app(home, brain)

    async def go():
        agent = app.state.agent
        t = agent.start("What time zone is Ghana in?")
        await asyncio.wait_for(t.runner, 10)
        agent.reply(t.id, "Is that the same as London?")
        await asyncio.wait_for(t.runner, 10)
        return t

    asyncio.run(go())
    second = brain.seen[1]
    assert [m["role"] for m in second] == ["user", "assistant", "user"]
    assert second[1]["content"] == "Ghana uses GMT."


def test_let_me_know_is_an_answer_not_an_unfinished_plan(home):
    said = "Paris is the capital of France. Let me know if you want me to look up more."
    brain = Scripted([said])
    task = asyncio.run(_run_plain(_app(home, brain), "capital of France?"))
    assert len(brain.seen) == 1  # no "you described what you would do" nudge
    assert task.status == "done" and _final(task) == said


def test_words_after_the_nudge_are_delivered_not_failed(home):
    brain = Scripted(["I'll check the prices for you.",
                      "I need to know which city you mean before I can look anything up."])
    task = asyncio.run(_run_plain(_app(home, brain), "find the best one"))
    assert task.status == "done"
    assert _final(task).startswith("I need to know which city")


def test_a_follow_up_about_a_file_saved_earlier_keeps_its_answer(home):
    target = home["downloads"] / "notes.txt"
    brain = Scripted([("", [("write_text_file", {"path": str(target), "content": "Notes."})]),
                      "It's saved in your Downloads folder.",
                      "It's saved in Downloads as notes.txt."])
    app = _app(home, brain)

    async def go():
        agent = app.state.agent
        t = agent.start("make a note of this")
        await asyncio.wait_for(t.runner, 10)
        agent.reply(t.id, "where did it go?")
        await asyncio.wait_for(t.runner, 10)
        return t

    task = asyncio.run(go())
    assert _final(task) == "It's saved in Downloads as notes.txt."


def test_pointing_is_never_backed_by_an_earlier_turn():
    assert unbacked_claims("I'm pointing at it now.", set(), {"point_at"}) == ["pointed at it"]
    assert unbacked_claims("It's saved in Documents.", set(), {"write_text_file"}) == []


def test_claudes_own_web_search_backs_i_searched(home):
    class SearchesServerSide(Scripted):
        async def step(self, system, messages, tools):
            self.seen.append(messages)
            blocks = [BetaServerToolUseBlock(id="s1", name="web_search", input={"query": "q"},
                                             type="server_tool_use"),
                      BetaTextBlock(text="I searched and the results show it opens at 9.", type="text")]
            return BrainReply(blocks, blocks[1].text, [], "end_turn")

    brain = SearchesServerSide([])
    task = asyncio.run(_run_plain(_app(home, brain), "when does the museum open?"))
    assert _final(task) == "I searched and the results show it opens at 9."


async def _run_plain(app, request):
    t = app.state.agent.start(request)
    await asyncio.wait_for(t.runner, 10)
    return t


def test_an_empty_reply_after_real_work_is_asked_for_words(home):
    target = home["downloads"] / "a.txt"
    brain = Scripted([("", [("write_text_file", {"path": str(target), "content": "x"})]), "",
                      "I wrote a.txt in Downloads."])
    task = asyncio.run(_run_plain(_app(home, brain), "write a file"))
    assert "reply was empty" in brain.seen[2][-1]["content"][0]["text"]
    assert task.status == "done" and _final(task) == "I wrote a.txt in Downloads."


def test_silence_after_real_work_reports_what_ran_not_nothing(home):
    target = home["downloads"] / "a.txt"
    brain = Scripted([("", [("write_text_file", {"path": str(target), "content": "x"})]), "", ""])
    task = asyncio.run(_run_plain(_app(home, brain), "write a file"))
    assert task.status == "done"
    assert _final(task).startswith("I stopped without saying whether this is finished, so check it. What ran:")
    assert target.exists()


def test_running_out_of_steps_ends_with_a_real_report(home):
    calls = [("", [("list_folder", {"path": str(home["downloads"])})])] * 3
    brain = Scripted([*calls, "I listed Downloads three times; the sorting is still left."])
    app = _app(home, brain)

    async def go():
        t = app.state.agent.start("sort my downloads")
        t.max_steps = 3
        await asyncio.wait_for(t.runner, 20)
        return t

    task = asyncio.run(go())
    assert "out of steps" in brain.seen[3][-1]["content"][0]["text"]
    assert _final(task) == "I listed Downloads three times; the sorting is still left. (I ran out of steps before finishing.)"


def test_a_dropped_connection_is_retried(home, monkeypatch):
    import compii_core.agent as agent_mod

    monkeypatch.setattr(agent_mod.asyncio, "sleep", _no_wait)

    class DropsOnce(Scripted):
        async def step(self, system, messages, tools):
            if not getattr(self, "dropped", False):
                self.dropped = True
                raise RuntimeError("RemoteProtocolError: Server disconnected without sending a response.")
            return await super().step(system, messages, tools)

    task = asyncio.run(_run_plain(_app(home, DropsOnce(["Hello!"])), "hi"))
    assert task.status == "done" and _final(task) == "Hello!"


_real_sleep = asyncio.sleep


async def _no_wait(seconds, *a, **kw):
    await _real_sleep(0)


def test_a_tool_that_keeps_failing_is_switched_off_for_the_turn(home):
    missing = str(home["downloads"] / "nope")
    calls = [("", [("list_folder", {"path": f"{missing}{i}"})]) for i in range(4)]
    brain = Scripted([*calls, "That folder isn't there."])
    task = asyncio.run(_run_plain(_app(home, brain), "list it"))
    results = [b for m in brain.seen[-1] if isinstance(m["content"], list) for b in m["content"]
               if isinstance(b, dict) and b.get("type") == "tool_result"]
    assert "failed 3 times in a row" in results[-1]["content"]
    ran = [e for e in task.events if e["type"] == "action_started" and e["tool"] == "list_folder"]
    assert len(ran) == 3


def test_web_tools_stop_when_the_pc_is_offline():
    from compii_core.agent import ToolCall, _FailureStreak

    streak = _FailureStreak()
    offline = {"is_error": True, "content": "ConnectError: [Errno 11001] getaddrinfo failed"}
    streak.saw(ToolCall("1", "read_web_page", {"url": "https://a"}), offline)
    streak.saw(ToolCall("2", "reach_platform", {"url": "https://b"}), offline)
    assert "can't reach the internet" in streak.refuse(ToolCall("3", "search_web", {"query": "x"}))
    assert streak.refuse(ToolCall("4", "list_folder", {"path": "C:/"})) == ""


def test_old_tool_results_shrink_for_small_models_only():
    from compii_core.agent import RECENT_RESULTS, shrink_old_results

    big = "x" * 15_000
    sent = [{"role": "user", "content": "go"}]
    for i in range(6):
        sent += [{"role": "assistant", "content": [{"type": "tool_use", "id": f"c{i}", "name": "t", "input": {}}]},
                 {"role": "user", "content": [{"type": "tool_result", "tool_use_id": f"c{i}", "content": big}]}]
    out = shrink_old_results(sent)
    sizes = [len(m["content"][0]["content"]) for m in out if m["role"] == "user" and isinstance(m["content"], list)]
    assert all(s < 1000 for s in sizes[:-RECENT_RESULTS]) and all(s == 15_000 for s in sizes[-RECENT_RESULTS:])
    assert len(sent[2]["content"][0]["content"]) == 15_000  # the conversation itself is untouched
    small = sent[:3]
    assert shrink_old_results(small) is small


def test_update_plan_takes_bare_titles_and_common_status_words(home):
    from compii_core.tools import Toolbox

    tb = Toolbox([home["downloads"]], home["workspace"])
    out = tb.execute("update_plan", {"steps": ["Make folders", {"title": "Move files", "status": "in_progress"}]})
    assert out.plan == [{"title": "Make folders", "status": "pending"}, {"title": "Move files", "status": "active"}]
    out = tb.execute("update_plan", {"steps": '[{"title": "Check", "status": "completed"}]'})
    assert out.plan == [{"title": "Check", "status": "done"}]


def test_an_empty_stream_from_the_gateway_is_asked_again(monkeypatch):
    import httpx

    from compii_core import providers

    monkeypatch.setattr(providers.asyncio, "sleep", _no_wait)
    sent = []

    def handler(request):
        sent.append(request)
        body = (b"data: [DONE]\n\n" if len(sent) == 1 else
                b'data: {"choices":[{"delta":{"content":"Sorted."},"finish_reason":"stop"}]}\n\ndata: [DONE]\n\n')
        return httpx.Response(200, headers={"content-type": "text/event-stream"}, content=body)

    brain = providers.OpenAICompatBrain("https://example.test/v1", "m", "k", transport=httpx.MockTransport(handler))

    async def noop(_):
        pass

    reply = asyncio.run(brain.step_stream("s", [{"role": "user", "content": "hi"}], [], noop))
    assert len(sent) == 2 and reply.text == "Sorted."


def test_an_empty_plain_reply_is_asked_again():
    import httpx

    from compii_core import providers

    sent = []

    def handler(request):
        sent.append(request)
        msg = {"content": ""} if len(sent) == 1 else {"content": "Hello."}
        return httpx.Response(200, json={"choices": [{"message": msg, "finish_reason": "stop"}]})

    brain = providers.OpenAICompatBrain("https://example.test/v1", "m", "k", transport=httpx.MockTransport(handler))
    reply = asyncio.run(brain.step("s", [{"role": "user", "content": "hi"}], []))
    assert len(sent) == 2 and reply.text == "Hello."


def test_compact_history_says_what_ran_and_drops_compii_notes():
    msgs = [
        {"role": "user", "content": "make notes.txt"},
        {"role": "assistant", "content": [{"type": "tool_use", "id": "a", "name": "write_text_file",
                                           "input": {"path": "C:/Users/me/Documents/notes.txt"}}]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "a", "content": "ok"},
                                     {"type": "text", "text": "(Check from Compii: your reply was empty.)"}]},
        {"role": "assistant", "content": [{"type": "text", "text": "Saved."}]},
    ]
    assert compact_history(msgs) == [
        {"role": "user", "content": "make notes.txt"},
        {"role": "assistant", "content": "(Compii ran: write_text_file notes.txt)\nSaved."}]


def test_a_copied_history_note_is_not_spoken(home):
    brain = Scripted(["Saved it. (Compii ran: write_text_file notes.txt)"])
    task = asyncio.run(_run_plain(_app(home, brain), "thanks"))
    assert _final(task) == "Saved it."


def test_move_labels_name_the_folder(home):
    from compii_core.tools import _move_label

    (home["downloads"] / "data").mkdir()
    d = home["downloads"]
    assert _move_label(str(d / "budget.csv"), str(d / "data" / "budget.csv")) == "Move budget.csv into data"
    assert _move_label(str(d / "budget.csv"), str(d / "data")) == "Move budget.csv into data"
    assert _move_label(str(d / "a.txt"), str(d / "b.txt")) == "Rename a.txt to b.txt"
    assert _move_label(str(d / "a.txt"), str(d / "data" / "b.txt")) == "Move a.txt to data/b.txt"


def test_a_job_report_names_pages_that_were_only_a_wall():
    from compii_core.agent import _unread_pages

    actions = [
        {"tool": "browse", "ok": True, "detail": "Page: Just a moment...\nAddress: https://www.macrotrends.net/x\n\nwww."},
        {"tool": "browse", "ok": True, "detail": "Page: 404 - File Not Found\nAddress: https://wpr.com/a\n\n404"},
        {"tool": "browse", "ok": True, "detail": "Page: Accra Population 2026\nAddress: https://wpr.com/b\n\nWORLD"},
        {"tool": "search_web", "ok": True, "detail": "1. Just a moment — https://x"},
    ]
    assert _unread_pages(actions) == ["https://www.macrotrends.net/x (Just a moment...)",
                                      "https://wpr.com/a (404 - File Not Found)"]


def test_a_rate_limited_page_counts_as_unread():
    from compii_core.agent import _unread_pages

    hit = {"tool": "browse", "ok": True,
           "detail": "Page: Wikimedia Error\nAddress: https://en.wikipedia.org/wiki/Asante_Empire\n\nError\nToo many requests."}
    assert _unread_pages([hit]) == ["https://en.wikipedia.org/wiki/Asante_Empire (Wikimedia Error)"]


def test_a_schedule_without_a_time_says_the_time_it_will_use():
    from compii_core.proactive import describe, validate_schedule

    spec = validate_schedule({"kind": "weekdays"})
    assert describe(spec) == "Every weekday at 09:00"
    assert validate_schedule({"kind": "daily", "at": "7:5"})["at"] == "07:05"


def test_a_spreadsheet_without_a_name_or_with_columns_still_lands(home):
    from openpyxl import load_workbook

    from compii_core.tools import ToolError, Toolbox

    tb = Toolbox([home["downloads"]], home["workspace"])
    p = home["downloads"] / "e.xlsx"
    tb.execute("create_spreadsheet", {"path": str(p), "sheets": {"columns": ["Item", "Amount"],
                                                                  "rows": [["Rent", 900], ["Total", "=SUM(B2:B2)"]]}})
    rows = list(load_workbook(p).active.iter_rows(values_only=True))
    assert rows == [("Item", "Amount"), ("Rent", 900), ("Total", "=SUM(B2:B2)")]
    try:
        tb.execute("create_spreadsheet", {"path": str(home["downloads"] / "f.xlsx"), "sheets": [{"name": "x"}]})
    except ToolError as e:
        assert "header row" in str(e)
    else:
        raise AssertionError("an empty sheet should be refused with a clear message")


def test_spreadsheet_cells_as_objects_and_rows_as_data(home):
    from openpyxl import load_workbook

    from compii_core.tools import Toolbox

    tb = Toolbox([home["downloads"]], home["workspace"])
    p = home["downloads"] / "g.xlsx"
    tb.execute("create_spreadsheet", {"path": str(p), "sheets": [{"name": "S", "data": [
        ["Item", "Amount"], ["rent", 900], ["Total", {"formula": "=SUM(B2:B2)"}]]}]})
    assert list(load_workbook(p).active.iter_rows(values_only=True)) == [
        ("Item", "Amount"), ("rent", 900), ("Total", "=SUM(B2:B2)")]
