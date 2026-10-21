"""Documents: Markdown that reads well, Word files that look made, and both read back again."""

import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest

from compii_core.documents import docx_bytes, docx_text, tidy_markdown
from compii_core.tools import Toolbox

MD = """# Scholarship shortlist

Two that fit the deadline.

| Name | Deadline | Award |
|---|---|---|
| Nexus | 12 Nov | $10,000 |
| Horizon | 5 Dec | $5,000 |

What I need:
- essays
- two references
  - one from a professor

1. Draft the essay
2. Upload the transcript

Read the [rules](https://example.com/nexus).
"""


@pytest.fixture
def home(tmp_path: Path) -> dict[str, Path]:
    downloads = tmp_path / "Downloads"
    downloads.mkdir()
    return {"downloads": downloads, "workspace": tmp_path / "Compii"}


@pytest.fixture
def toolbox(home) -> Toolbox:
    return Toolbox([home["downloads"]], home["workspace"], lambda s: None)


# ---- Markdown -------------------------------------------------------------------


def test_headings_bullets_and_blank_lines_are_normalised():
    messy = "Title\n===\n\n\n#  Heading  \n* item\n+ item2\nbody\n"
    assert tidy_markdown(messy) == "# Title\n\n# Heading\n\n- item\n- item2\n\nbody\n"


def test_tables_line_up_and_a_stray_bold_marker_goes():
    out = tidy_markdown("| A|B |\n|:--|--:|\n|one|two|\n\nBold ** text\n")
    assert "| A | B |" in out and "| one | two |" in out
    assert "**" not in out, "an unclosed ** prints as two stars in Word"


def test_a_paragraph_under_a_list_is_not_swallowed_by_it():
    assert tidy_markdown("- first\nsecond line\n") == "- first\n\nsecond line\n"


def test_code_blocks_keep_every_line():
    out = tidy_markdown("```python\nx = 1\n\n  y = 2   \n```\n")
    assert "x = 1" in out and "  y = 2" in out


# ---- Word -----------------------------------------------------------------------


def test_the_word_package_is_well_formed_and_carries_the_formatting(tmp_path: Path):
    path = tmp_path / "shortlist.docx"
    path.write_bytes(docx_bytes(MD, title="Scholarship shortlist"))
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            ET.fromstring(z.read(name))  # nothing in the package is broken XML
        doc = z.read("word/document.xml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
        core = z.read("docProps/core.xml").decode()
    assert '<w:pStyle w:val="Heading1"/>' in doc
    assert "<w:tbl>" in doc and 'w:fill="DCE6F1"' in doc, "the header row is shaded"
    assert '<w:numId w:val="1"/>' in doc and '<w:numId w:val="2"/>' in doc, "real bullets and numbers"
    assert "https://example.com/nexus" in rels and "<w:hyperlink" in doc
    assert "<dc:title>Scholarship shortlist</dc:title>" in core


def test_every_link_the_document_uses_is_declared_in_the_package(tmp_path: Path):
    """A reference to a relationship that isn't there is what makes Word say a file is unreadable."""
    path = tmp_path / "links.docx"
    path.write_bytes(docx_bytes("# A\n\nSee [one](https://a.example/1) and [two](https://b.example/2).\n"))
    with zipfile.ZipFile(path) as z:
        used = set(re.findall(r'r:id="(r\d+)"', z.read("word/document.xml").decode()))
        declared = set(re.findall(r'Relationship Id="(r\d+)"', z.read("word/_rels/document.xml.rels").decode()))
    assert used and used <= declared
    assert len(used) == 2


def test_lines_written_one_under_the_other_stay_that_way():
    text = docx_text_of("Dear committee,\nAdmissions Office\nRoom 14\n")
    assert "Dear committee,\nAdmissions Office\nRoom 14" in text


def test_compii_reads_its_own_word_file_back_as_text():
    text = docx_text_of(MD)
    assert "# Scholarship shortlist" in text
    assert "| Nexus | 12 Nov | $10,000 |" in text
    assert "- essays" in text and "two references" in text


def docx_text_of(markdown: str) -> str:
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "x.docx"
        p.write_bytes(docx_bytes(markdown))
        return docx_text(p)


# ---- the tools Compii calls -----------------------------------------------------


def test_create_document_always_lands_a_docx_that_reads_back(toolbox, home):
    res = toolbox.execute("create_document", {"path": "shortlist", "content": MD, "title": "Scholarship shortlist"})
    docx = home["workspace"] / "shortlist.docx"
    assert docx.is_file(), res.text
    assert "Scholarship shortlist" in toolbox.execute("read_file", {"path": str(docx)}).text
    assert (str(docx), "created") in res.files


def test_create_document_refuses_to_replace_a_file(toolbox, home):
    toolbox.execute("create_document", {"path": "letter.docx", "content": "# Hello"})
    with pytest.raises(Exception, match="already exists"):
        toolbox.execute("create_document", {"path": "letter.docx", "content": "# Twice"})


def test_markdown_written_by_the_plain_file_tool_is_tidied(toolbox, home):
    toolbox.execute("write_text_file", {"path": "notes.md", "content": "# Title\n* one\n+ two\n"})
    assert (home["workspace"] / "notes.md").read_text(encoding="utf-8") == "# Title\n\n- one\n- two\n"


def test_spreadsheets_get_a_frozen_filterable_header(toolbox, home):
    from openpyxl import load_workbook

    toolbox.execute("create_spreadsheet", {"path": "t.xlsx", "sheets": [
        {"name": "S", "header": ["Item", "Amount"], "rows": [["a", 1], ["b", 2]]}]})
    ws = load_workbook(home["workspace"] / "t.xlsx")["S"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions
    assert ws["A1"].font.bold and str(ws["A1"].fill.fgColor.rgb).endswith("1F3864")
